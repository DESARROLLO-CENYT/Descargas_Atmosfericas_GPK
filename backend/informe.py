"""Generacion del informe en Excel de la vista de datos.

Vive aparte de main.py porque es codigo de presentacion, no de calculo: recibe
la respuesta ya armada de /api/procesar y la vuelca en un libro con formato.
Asi el analisis tiene una sola fuente de verdad y este modulo no repite reglas.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta alineada con el tablero. El fondo va claro y no oscuro: el informe se
# imprime y se reenvia por correo, donde el tema oscuro gasta tinta y se lee mal.
ROJO = "C90016"          # acento de marca, encabezados
GRIS_TITULO = "1A1A1F"
GRIS_SUAVE = "F4F4F6"    # bandeado de filas
GRIS_BORDE = "D9D9DE"
MORADO = "7E22CE"        # DPS/DSD, igual que en los mapas
AMBAR = "B45309"         # error de localizacion que supera la distancia
VERDE = "15803D"

BORDE = Border(*[Side(style="thin", color=GRIS_BORDE)] * 4)


def _stats_corriente(rayos):
    """Media, mediana y percentil 99 de la magnitud de corriente."""
    vals = sorted(abs(float(r.get("corriente") or 0)) for r in rayos)
    vals = [v for v in vals if v > 0]
    if not vals:
        return None

    def pct(p):
        if len(vals) == 1:
            return vals[0]
        pos = (len(vals) - 1) * p
        bajo, alto = int(pos), min(int(pos) + 1, len(vals) - 1)
        return vals[bajo] + (pos - bajo) * (vals[alto] - vals[bajo])

    return {
        "n": len(vals),
        "media": sum(vals) / len(vals),
        "mediana": pct(0.5),
        "p99": pct(0.99),
        "min": vals[0],
        "max": vals[-1],
    }


def _agrupar(estructuras, clave, extras):
    """Mismo agregado que hacen las tablas de resumen del frontend."""
    grupos = {}
    for e in estructuras:
        k = clave(e)
        if k not in grupos:
            grupos[k] = dict(extras(e), estructuras=0, afectadas=0, impactos=0, corriente_max=0.0)
        g = grupos[k]
        g["estructuras"] += 1
        if (e.get("impactos") or 0) > 0:
            g["afectadas"] += 1
        g["impactos"] += e.get("impactos") or 0
        if abs(e.get("corriente_max") or 0) > abs(g["corriente_max"]):
            g["corriente_max"] = e.get("corriente_max") or 0
    return list(grupos.values())


def _ancho(columnas, filas):
    """Ancho por columna segun su contenido mas largo, acotado a un rango sano.

    Es lo que hace que las columnas queden proporcionales: 'TAG' no ocupa lo
    mismo que 'Pórtico / Circuito', que llega a 40 caracteres.
    """
    anchos = []
    for col in columnas:
        largo = len(str(col["titulo"]))
        for f in filas:
            v = f.get(col["clave"])
            if v is None:
                continue
            texto = f"{v:,.1f}" if isinstance(v, float) else str(v)
            largo = max(largo, len(texto))
        anchos.append(min(46, max(10, largo + 3)))
    return anchos


def _hoja_tabla(wb, titulo, subtitulo, columnas, filas):
    """Escribe una hoja con encabezado de marca, banda de filas y autofiltro."""
    ws = wb.create_sheet(titulo)
    ws.sheet_view.showGridLines = False

    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=GRIS_TITULO)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(size=9, color="666670")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columnas))
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 6

    fila_cab = 4
    for i, col in enumerate(columnas, start=1):
        c = ws.cell(row=fila_cab, column=i, value=col["titulo"])
        c.fill = PatternFill("solid", fgColor=ROJO)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila_cab].height = 30

    for j, fila in enumerate(filas):
        r = fila_cab + 1 + j
        banda = PatternFill("solid", fgColor=GRIS_SUAVE) if j % 2 else None
        for i, col in enumerate(columnas, start=1):
            v = fila.get(col["clave"])
            c = ws.cell(row=r, column=i)
            tipo = col.get("tipo")

            if tipo == "bool":
                c.value = "Sí" if v else "No"
                c.font = Font(bold=bool(v), color=MORADO if v else "8A8A93", size=10)
                c.alignment = Alignment(horizontal="center")
            else:
                c.value = v
                c.font = Font(size=10)
                if col.get("formato"):
                    c.number_format = col["formato"]
                    c.alignment = Alignment(horizontal="right")
                # El error que supera a la distancia se marca: ahi el valor de
                # la distancia deja de poder afirmarse
                if tipo == "error" and v is not None and (fila.get("dist_min") or 0) < v:
                    c.font = Font(size=10, bold=True, color=AMBAR)

            if banda:
                c.fill = banda
            c.border = BORDE

    for i, ancho in enumerate(_ancho(columnas, filas), start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    if filas:
        ws.auto_filter.ref = (
            f"A{fila_cab}:{get_column_letter(len(columnas))}{fila_cab + len(filas)}"
        )
    ws.freeze_panes = ws.cell(row=fila_cab + 1, column=1)
    return ws


def _hoja_resumen(wb, datos, filtros, stats):
    ws = wb.create_sheet("Resumen", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30

    ws["A1"] = "Informe de Descargas Atmosféricas"
    ws["A1"].font = Font(bold=True, size=18, color=GRIS_TITULO)
    ws["A2"] = "Campo Llanos 34 · GeoPark"
    ws["A2"].font = Font(size=11, color=ROJO, bold=True)
    ws["A3"] = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}"
    ws["A3"].font = Font(size=9, color="666670")

    fila = 5

    def seccion(titulo):
        nonlocal fila
        c = ws.cell(row=fila, column=1, value=titulo)
        c.fill = PatternFill("solid", fgColor=ROJO)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        ws.cell(row=fila, column=2).fill = PatternFill("solid", fgColor=ROJO)
        ws.row_dimensions[fila].height = 20
        fila += 1

    def dato(etiqueta, valor, formato=None, color=None):
        nonlocal fila
        a = ws.cell(row=fila, column=1, value=etiqueta)
        a.font = Font(size=10, color="44444C")
        b = ws.cell(row=fila, column=2, value=valor)
        b.font = Font(size=10, bold=True, color=color or GRIS_TITULO)
        if formato:
            b.number_format = formato
        a.border = BORDE
        b.border = BORDE
        fila += 1

    k = datos.get("kpis", {})
    seccion("Filtros aplicados")
    for etiqueta, valor in filtros:
        dato(etiqueta, valor)

    fila += 1
    seccion("Resultados")
    dato("Estructuras analizadas", k.get("total_estructuras", 0), "#,##0")
    dato("Estructuras afectadas", k.get("estructuras_afectadas", 0), "#,##0", ROJO)
    dato(f"Descargas dentro del radio ({k.get('radio', 0):.0f} m)", k.get("total_rayos", 0), "#,##0")
    dato("Descargas en el rango de fechas", k.get("total_rayos_rango", 0), "#,##0")
    rango = k.get("total_rayos_rango") or 0
    dato("Tasa de exposición", (k.get("total_rayos", 0) / rango) if rango else 0, "0.00%")

    if stats:
        fila += 1
        seccion("Corriente de las descargas (magnitud)")
        dato("Descargas consideradas", stats["n"], "#,##0")
        dato("Media", stats["media"], '#,##0.0 "kA"')
        dato("Mediana", stats["mediana"], '#,##0.0 "kA"')
        dato("Percentil 99", stats["p99"], '#,##0.0 "kA"', ROJO)
        dato("Mínima", stats["min"], '#,##0.0 "kA"')
        dato("Máxima", stats["max"], '#,##0.0 "kA"')

    fila += 1
    seccion("Nota sobre la distancia mínima")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 2, end_column=2)
    aviso = ws.cell(
        row=fila, column=1,
        value=("La distancia al rayo más cercano se acompaña del error de localización "
               "de esa misma descarga. Cuando el error supera a la distancia (celda en "
               "ámbar), la cercanía no puede afirmarse: el rayo pudo caer bastante más "
               "lejos de lo que indica el valor."),
    )
    aviso.alignment = Alignment(wrap_text=True, vertical="top")
    aviso.font = Font(size=9, color="44444C")
    ws.row_dimensions[fila].height = 46
    return ws


def construir_informe(datos, filtros):
    """Arma el libro completo y lo devuelve como bytes listos para descargar."""
    wb = Workbook()
    wb.remove(wb.active)

    estructuras = datos.get("estructuras", [])
    rayos = datos.get("rayos", [])
    impactadas = sorted(
        [e for e in estructuras if (e.get("impactos") or 0) > 0],
        key=lambda e: -(e.get("impactos") or 0),
    )

    stats = _stats_corriente(rayos)
    _hoja_resumen(wb, datos, filtros, stats)

    _hoja_tabla(
        wb, "Estructuras",
        "Solo estructuras impactadas, de mayor a menor número de impactos.",
        [
            {"clave": "id", "titulo": "TAG"},
            {"clave": "campo", "titulo": "Campo"},
            {"clave": "locacion", "titulo": "Locación"},
            {"clave": "portico", "titulo": "Pórtico / Circuito"},
            {"clave": "impactos", "titulo": "Impactos", "formato": "#,##0"},
            {"clave": "corriente_max", "titulo": "Corriente máx (kA)", "formato": "#,##0.0"},
            {"clave": "dist_min", "titulo": "Dist. mínima (m)", "formato": "#,##0"},
            {"clave": "error_min", "titulo": "± error (m)", "formato": "#,##0", "tipo": "error"},
            {"clave": "dps", "titulo": "DPS", "tipo": "bool"},
            {"clave": "dsd", "titulo": "DSD", "tipo": "bool"},
            {"clave": "lat", "titulo": "Latitud", "formato": "0.000000"},
            {"clave": "lon", "titulo": "Longitud", "formato": "0.000000"},
        ],
        impactadas,
    )

    _hoja_tabla(
        wb, "Rayos",
        "Descargas dentro del radio de búsqueda, de la más reciente a la más antigua.",
        [
            {"clave": "fecha", "titulo": "Fecha"},
            {"clave": "corriente", "titulo": "Corriente (kA)", "formato": "#,##0.0"},
            {"clave": "magnitud", "titulo": "Magnitud (kA)", "formato": "#,##0.0"},
            {"clave": "polaridad", "titulo": "Polaridad"},
            {"clave": "lat", "titulo": "Latitud", "formato": "0.0000"},
            {"clave": "lon", "titulo": "Longitud", "formato": "0.0000"},
        ],
        sorted(
            [dict(r, magnitud=abs(float(r.get("corriente") or 0))) for r in rayos],
            key=lambda r: str(r.get("fecha") or ""), reverse=True,
        ),
    )

    por_circuito = [g for g in _agrupar(
        estructuras, lambda e: e.get("portico"),
        lambda e: {"portico": e.get("portico") or "(sin asignar)",
                   "campo": e.get("campo") or "—",
                   "locacion": e.get("locacion") or "—"},
    ) if g["impactos"] > 0]

    _hoja_tabla(
        wb, "Por circuito",
        "Solo circuitos con impactos, de mayor a menor.",
        [
            {"clave": "portico", "titulo": "Pórtico / Circuito"},
            {"clave": "campo", "titulo": "Campo"},
            {"clave": "locacion", "titulo": "Locación"},
            {"clave": "estructuras", "titulo": "Estructuras", "formato": "#,##0"},
            {"clave": "afectadas", "titulo": "Afectadas", "formato": "#,##0"},
            {"clave": "impactos", "titulo": "Impactos", "formato": "#,##0"},
            {"clave": "corriente_max", "titulo": "Corriente máx (kA)", "formato": "#,##0.0"},
        ],
        sorted(por_circuito, key=lambda g: -g["impactos"]),
    )

    por_campo = [g for g in _agrupar(
        estructuras, lambda e: (e.get("campo"), e.get("locacion")),
        lambda e: {"campo": e.get("campo") or "(sin asignar)",
                   "locacion": e.get("locacion") or "—"},
    ) if g["impactos"] > 0]

    _hoja_tabla(
        wb, "Por campo",
        "Solo locaciones con impactos, de mayor a menor.",
        [
            {"clave": "campo", "titulo": "Campo"},
            {"clave": "locacion", "titulo": "Locación"},
            {"clave": "estructuras", "titulo": "Estructuras", "formato": "#,##0"},
            {"clave": "afectadas", "titulo": "Afectadas", "formato": "#,##0"},
            {"clave": "impactos", "titulo": "Impactos", "formato": "#,##0"},
            {"clave": "corriente_max", "titulo": "Corriente máx (kA)", "formato": "#,##0.0"},
        ],
        sorted(por_campo, key=lambda g: -g["impactos"]),
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
